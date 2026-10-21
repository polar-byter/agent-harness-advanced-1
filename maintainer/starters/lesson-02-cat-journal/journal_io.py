"""Чтение и запись журнала в Excel. Модель здесь не используется."""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADER = ['Дата', 'Кто записал', 'Что случилось']
RESULT_SHEET = 'Разбор'


def as_text(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime('%d.%m.%Y')
    return '' if value is None else str(value).strip()


def read_journal(path: Path) -> list[dict]:
    """Первый лист: дата, автор записи и текст. Строки без текста пропускаем."""
    book = load_workbook(path, read_only=True)
    sheet = book.worksheets[0]
    entries = []
    for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        when, author, text = (as_text(value) for value in row)
        if text:
            entries.append({'date': when, 'author': author, 'text': text})
    book.close()
    return entries


def save_results(path: Path, results: list[dict], summary: str) -> None:
    """Пересоздаём лист с разбором, исходный лист журнала не трогаем."""
    book = load_workbook(path)
    if RESULT_SHEET in book.sheetnames:
        del book[RESULT_SHEET]
    sheet = book.create_sheet(RESULT_SHEET)
    sheet.append(HEADER + ['Категория', 'Серьёзность'])
    for item in results:
        sheet.append([item['date'], item['author'], item['text'], item['category'], item['severity']])
    sheet.append([])
    sheet.append(['Сводка для семьи', '', summary])
    for column, width in zip('ABCDE', (12, 14, 60, 12, 12)):
        sheet.column_dimensions[column].width = width
    book.save(path)


def append_entry(path: Path, author: str, text: str, extra: tuple = ()) -> None:
    """Добавляем запись в конец журнала; если файла нет, создаём его с заголовком.

    В extra можно сразу записать категорию и серьёзность: они попадут в столбцы D и E.
    """
    if path.is_file():
        book = load_workbook(path)
    else:
        book = Workbook()
        book.active.title = 'Журнал'
        book.active.append(HEADER)
    book.worksheets[0].append([date.today().strftime('%d.%m.%Y'), author, text, *extra])
    book.save(path)


def read_classified(path: Path) -> list[dict]:
    """Записи, у которых в столбцах D и E уже есть категория и серьёзность."""
    book = load_workbook(path, read_only=True)
    entries = []
    for row in book.worksheets[0].iter_rows(min_row=2, max_col=5, values_only=True):
        when, author, text, category, severity = (list(row) + [None] * 5)[:5]
        if text and category:
            entries.append({'date': as_text(when), 'author': as_text(author), 'text': as_text(text),
                            'category': as_text(category), 'severity': severity})
    book.close()
    return entries
